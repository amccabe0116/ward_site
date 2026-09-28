#!/usr/bin/env python3
"""Convert the two leadership Google Sheets (exported as .xlsx) into the JSON that
admin_replace_sheet() stores — the same shape scripts/announcements.gs (syncMemberSheets)
sends. Manual fallback when the Apps Script isn't set up:

    python3 scripts/sheets_to_json.py callings.xlsx newmember.xlsx > sheets.json

Output: {"callings": {...}, "callings_committees": {...}, "newmember": {...}} where each is
{title, source_url, headers:[…], rows:[[…], …]} with every cell a string ("" when blank).
"""
import datetime, json, sys
import openpyxl

CALLINGS_URL = 'https://docs.google.com/spreadsheets/d/1PMS3f4ncGaeIhJJ9ZaVAbOgA0kUTBnMOWgpvhKgMbe8/edit'
NEWMEMBER_URL = 'https://docs.google.com/spreadsheets/d/1OyPHy_STcN-Nbh_OiPVIiSIu16cq1gxvs1ffzgePLlA/edit?gid=1364401202#gid=1364401202'


def cell(v):
    if v is None:
        return ''
    if isinstance(v, datetime.datetime):
        if v.hour or v.minute or v.second:
            return f'{v.month}/{v.day}/{v.year} {v.hour}:{v.minute:02d}:{v.second:02d}'
        return f'{v.day} {v.strftime("%b")} {v.year}'
    if isinstance(v, datetime.date):
        return f'{v.day} {v.strftime("%b")} {v.year}'
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def sheet_json(ws, title, url):
    rows = list(ws.iter_rows(values_only=True))
    headers = [cell(h) for h in rows[0]]
    while headers and headers[-1] == '':
        headers.pop()
    n = len(headers)
    out = []
    for r in rows[1:]:
        vals = [cell(v) for v in r[:n]] + [''] * max(0, n - len(r))
        if any(vals):
            out.append(vals)
    return {'title': title, 'source_url': url, 'headers': headers, 'rows': out}


def main(callings_path, newmember_path):
    cw = openpyxl.load_workbook(callings_path, data_only=True)
    nw = openpyxl.load_workbook(newmember_path, data_only=True)
    out = {
        'callings': sheet_json(cw.worksheets[0], 'Members without callings', CALLINGS_URL),
        'newmember': sheet_json(nw.worksheets[0], 'New member form', NEWMEMBER_URL),
    }
    if len(cw.worksheets) > 1:
        out['callings_committees'] = sheet_json(cw.worksheets[1], 'Committee requests', CALLINGS_URL)
    json.dump(out, sys.stdout, ensure_ascii=False)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
